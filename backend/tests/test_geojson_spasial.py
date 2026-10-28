"""Unggah poligon GeoJSON + query spasial + ekspor. Poligon uji dibuat DI LUAR Kab. Bandung (sudut timur laut kotak validasi),
sehingga tidak bentrok dengan poligon asli. Dibersihkan otomatis."""
import csv
import io
import json

import psycopg
import pytest
from openpyxl import load_workbook

from app.db import conninfo

NAMA = "Uji Spasial"


def kotak(x0, y0, x1, y1):
    return {"type": "Polygon", "coordinates": [[[x0, y0], [x1, y0], [x1, y1], [x0, y1], [x0, y0]]]}


KOTAK_A = kotak(108.00, -6.75, 108.10, -6.65)            # ~122 km2, titik tengah (108.05, -6.70)
KOTAK_B = kotak(108.05, -6.75, 108.15, -6.65)            # tumpang tindih dengan A


def geojson(*fitur):
    return json.dumps({"type": "FeatureCollection", "features": [
        {"type": "Feature", "properties": p, "geometry": g} for p, g in fitur]}).encode()


def unggah(client, header, isi, dry_run=False):
    return client.post("/api/polsek/geojson", headers=header["admin"], data={"dry_run": str(dry_run).lower()},
                       files={"file": ("poligon.geojson", isi, "application/geo+json")})


@pytest.fixture()
def polsek_uji(client, header, db_hidup):
    def bersihkan():
        with psycopg.connect(conninfo(), autocommit=True) as c:
            c.execute("DELETE FROM polsek WHERE nama_polsek IN (%s, 'Uji Spasial Dua')", (NAMA,))
    bersihkan()
    p = client.post("/api/polsek", json={"nama_polsek": NAMA}, headers=header["admin"]).json()
    yield p["id_polsek"]
    bersihkan()


def test_unggah_hanya_admin(client, header):
    assert client.post("/api/polsek/geojson", files={"file": ("a.geojson", b"{}")}).status_code == 401
    r = client.post("/api/polsek/geojson", headers=header["operator"], files={"file": ("a.geojson", b"{}")})
    assert r.status_code == 403


def test_dry_run_lalu_simpan(client, header, polsek_uji):
    isi = geojson(({"id_polsek": polsek_uji, "nama_polsek": NAMA}, KOTAK_A))
    r = unggah(client, header, isi, dry_run=True)
    assert r.status_code == 200 and r.json()["disimpan"] is False and r.json()["polsek_diperbarui"][0]["id_polsek"] == polsek_uji
    with psycopg.connect(conninfo()) as c:
        assert c.execute("SELECT geom IS NULL FROM polsek WHERE id_polsek = %s", (polsek_uji,)).fetchone()[0] is True   # belum tersimpan
    r = unggah(client, header, isi)
    assert r.status_code == 200 and r.json()["disimpan"] is True
    with psycopg.connect(conninfo()) as c:
        assert c.execute("SELECT GeometryType(geom), ST_SRID(geom) FROM polsek WHERE id_polsek = %s", (polsek_uji,)).fetchone() == ("MULTIPOLYGON", 4326)


def test_cocok_lewat_nama_tanpa_id(client, header, polsek_uji):
    r = unggah(client, header, geojson(({"nama_polsek": "POLSEK UJI SPASIAL"}, KOTAK_A)))
    assert r.status_code == 200 and r.json()["polsek_diperbarui"] == [{"id_polsek": polsek_uji, "nama_polsek": NAMA}]


@pytest.mark.parametrize("props, geom, potongan", [
    ({"id_polsek": 9999999}, KOTAK_A, "tidak ada di database"),
    ({"nama_polsek": "Tidak Dikenal Sama Sekali"}, KOTAK_A, "tidak dikenal"),
    ({}, KOTAK_A, "properties.id_polsek"),
    ({"id_polsek": "X"}, KOTAK_A, "bukan bilangan bulat"),
    ({"id_polsek": "ID", "nama_polsek": "Nama Lain"}, KOTAK_A, "tertukar"),
    ({"id_polsek": "ID"}, {"type": "Point", "coordinates": [108.05, -6.7]}, "Polygon atau MultiPolygon"),
    ({"id_polsek": "ID"}, kotak(-6.75, 108.00, -6.65, 108.10), "di luar kotak wilayah"),      # lon/lat tertukar
    ({"id_polsek": "ID"}, kotak(108.00, -6.7500, 108.0001, -6.7499), "tidak wajar"),            # terlalu kecil
])
def test_geojson_tidak_valid(client, header, polsek_uji, props, geom, potongan):
    props = {k: (polsek_uji if v == "ID" else v) for k, v in props.items()}
    r = unggah(client, header, geojson((props, geom)))
    assert r.status_code == 400, r.text
    assert potongan in " ".join(r.json()["error"]["detail"])


def test_berkas_rusak_dan_tumpang_tindih(client, header, polsek_uji):
    assert unggah(client, header, b"bukan json").status_code == 400
    assert unggah(client, header, b'{"type": "Point"}').status_code == 400
    assert unggah(client, header, geojson()).status_code == 400
    assert unggah(client, header, geojson(({"id_polsek": polsek_uji}, KOTAK_A))).status_code == 200
    kedua = client.post("/api/polsek", json={"nama_polsek": "Uji Spasial Dua"}, headers=header["admin"]).json()["id_polsek"]
    r = unggah(client, header, geojson(({"id_polsek": kedua}, KOTAK_B)))
    assert r.status_code == 400 and "tumpang tindih" in " ".join(r.json()["error"]["detail"])
    with psycopg.connect(conninfo()) as c:                                                       # gagal -> tidak ada yang berubah
        assert c.execute("SELECT geom IS NULL FROM polsek WHERE id_polsek = %s", (kedua,)).fetchone()[0] is True


def test_ganti_poligon_satu_polsek(client, header, polsek_uji):
    r = client.put(f"/api/polsek/{polsek_uji}/poligon", json={"type": "Feature", "properties": {}, "geometry": KOTAK_A}, headers=header["admin"])
    assert r.status_code == 200 and r.json()["disimpan"] is True
    assert client.put("/api/polsek/9999999/poligon", json=KOTAK_A, headers=header["admin"]).status_code == 404
    assert client.put(f"/api/polsek/{polsek_uji}/poligon", json={"type": "Point", "coordinates": [1, 2]}, headers=header["admin"]).status_code == 400


def test_spasial(client, header, polsek_uji, bersih_tahun_uji):
    unggah(client, header, geojson(({"id_polsek": polsek_uji}, KOTAK_A)))
    client.post("/api/kasus", json={"id_polsek": polsek_uji, "id_jenis": 1, "tahun": 2099, "jumlah": 45}, headers=header["operator"])

    r = client.get("/api/spasial/titik?lat=-6.70&lng=108.05&tahun=2099").json()                  # di dalam
    assert r["di_dalam_wilayah"] is True and r["polsek"]["id_polsek"] == polsek_uji
    assert (r["polsek"]["total_kasus"], r["polsek"]["kategori_zona"], r["polsek"]["jarak_meter"]) == (45, "sedang", 0)

    r = client.get("/api/spasial/titik?lat=-6.70&lng=108.11&tahun=2099").json()                  # ~1 km di luar -> Polsek terdekat
    assert r["di_dalam_wilayah"] is False and r["polsek"]["id_polsek"] == polsek_uji and 500 < r["polsek"]["jarak_meter"] < 2000

    t = client.get("/api/spasial/terdekat?lat=-6.70&lng=108.11&limit=3&radius_m=5000&tahun=2099").json()
    assert [x["id_polsek"] for x in t["data"]] == [polsek_uji]                                   # hanya yang dalam radius 5 km
    assert client.get("/api/spasial/terdekat?lat=-6.70&lng=108.11&radius_m=100&tahun=2099").json()["data"] == []

    b = client.get("/api/spasial/bbox?min_lng=108.04&min_lat=-6.72&max_lng=108.06&max_lat=-6.68&tahun=2099").json()
    assert [f["properties"]["id_polsek"] for f in b["features"]] == [polsek_uji] and b["features"][0]["geometry"]["type"] == "MultiPolygon"
    b = client.get("/api/spasial/bbox?min_lng=108.04&min_lat=-6.72&max_lng=108.06&max_lat=-6.68&tahun=2099&geometri=false").json()
    assert b["features"][0]["geometry"] is None
    assert client.get("/api/spasial/bbox?min_lng=108.04&min_lat=-6.72&max_lng=108.06&max_lat=-6.68&tahun=2099&zona=rendah").json()["features"] == []


@pytest.mark.parametrize("url", ["/api/spasial/titik?lat=95&lng=108", "/api/spasial/titik?lat=1", "/api/spasial/terdekat?lat=1&lng=1&limit=0",
                                 "/api/spasial/bbox?min_lng=2&min_lat=0&max_lng=1&max_lat=1"])
def test_spasial_validasi(client, db_hidup, url):
    assert client.get(url).status_code == 400


def test_ekspor_csv_xlsx(client, header, bersih_tahun_uji):
    client.post("/api/kasus", json={"id_polsek": 1, "id_jenis": 1, "tahun": 2099, "jumlah": 61}, headers=header["operator"])
    nama1 = client.get("/api/polsek/1?tahun=2099").json()["nama_polsek"]

    r = client.get("/api/export?format=csv&tahun=2099")
    assert r.status_code == 200 and r.headers["content-type"].startswith("text/csv")
    assert r.headers["content-disposition"] == 'attachment; filename="rekap_kriminalitas_2099.csv"'
    rows = list(csv.reader(io.StringIO(r.content.decode("utf-8-sig"))))
    assert rows[0][0] == "Polsek" and rows[0][1] == "Aniaya" and rows[0][-2:] == ["Total", "Zona"] and len(rows[0]) == 1 + 18 + 2
    baris = {x[0]: x for x in rows[1:]}[nama1]
    assert baris[1] == "61" and baris[-2] == "61" and baris[-1] == "tinggi"

    r = client.get("/api/export?format=csv&tahun=2099&jenis=1&polsek=1")                         # filter: 1 kolom jenis, 1 polsek
    rows = list(csv.reader(io.StringIO(r.content.decode("utf-8-sig"))))
    assert rows[0] == ["Polsek", "Aniaya", "Total", "Zona"] and rows[1] == [nama1, "61", "61", "tinggi"] and len(rows) == 2

    r = client.get("/api/export?format=xlsx&tahun=2099")
    assert r.status_code == 200 and r.headers["content-disposition"].endswith('.xlsx"')
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws.title == "Rekap 2099" and ws["A1"].value == "Polsek" and ws.cell(row=ws.max_row, column=1).value == "TOTAL"
    assert ws.cell(row=ws.max_row, column=2).value == 61

    assert client.get("/api/export?format=pdf").status_code == 400
    assert client.get("/api/export?tahun=2097").status_code == 404
