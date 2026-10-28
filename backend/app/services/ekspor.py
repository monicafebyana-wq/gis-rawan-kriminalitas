"""Susun & tulis tabel ekspor rekap (CSV / XLSX). Format kolom dibuat sama dengan impor CSV:
Polsek + kolom jenis (nama kolom baku, mis. 'Curanmor R2') + 'Total' + 'Zona'. 'Total' & 'Zona' diabaikan saat impor ulang.
'Zona' selalu dari total SEMUA jenis (sama seperti peta), walau kolom jenis difilter."""
import csv
import io

import psycopg

from app.services import rekap
from app.services.impor_csv import KOLOM_JENIS

XLSX_MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_WARNA_ZONA = {"tinggi": "F8B4B4", "sedang": "FCE8A6", "rendah": "B7E4C7"}


def susun_tabel(conn: psycopg.Connection, tahun: int, jenis: list[int] | None, polsek: list[int] | None) -> tuple[list[str], list[list]]:
    daftar_jenis = conn.execute(
        """SELECT id_jenis, kode_jenis, nama_jenis FROM jenis_kejahatan
           WHERE (%(j)s::int[] IS NULL OR id_jenis = ANY(%(j)s::int[])) ORDER BY id_jenis""", {"j": jenis}).fetchall()
    judul_baku = {kode: kolom for kolom, kode in KOLOM_JENIS}
    header = ["Polsek", *[judul_baku.get(j["kode_jenis"], j["nama_jenis"]) for j in daftar_jenis], "Total", "Zona"]
    sel = {(r["id_polsek"], r["id_jenis"]): r["jumlah"]
           for r in conn.execute("SELECT id_polsek, id_jenis, jumlah FROM kasus_kejahatan WHERE tahun = %s", (tahun,)).fetchall()}
    baris = []
    for r in rekap.rekap_per_polsek(conn, tahun, jenis, polsek):
        nilai = [sel.get((r["id_polsek"], j["id_jenis"]), 0) for j in daftar_jenis]
        baris.append([r["nama_polsek"], *nilai, sum(nilai), r["kategori_zona"]])
    return header, baris


def ke_csv(header: list[str], baris: list[list]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(header)
    w.writerows(baris)
    return ("\ufeff" + buf.getvalue()).encode("utf-8")          # BOM agar Excel membaca UTF-8 dengan benar


def ke_xlsx(header: list[str], baris: list[list], tahun: int) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = f"Rekap {tahun}"
    ws.append(header)
    for b in baris:
        ws.append(b)
    if baris:                                                    # baris total (hanya di XLSX; CSV dibiarkan bersih untuk impor ulang)
        ws.append(["TOTAL", *[sum(b[i] for b in baris) for i in range(1, len(header) - 2)],
                   sum(b[-2] for b in baris), ""])
        for c in ws[ws.max_row]:
            c.font = Font(bold=True)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3A5F")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    kol_zona = len(header)
    for r in range(2, 2 + len(baris)):
        sel = ws.cell(row=r, column=kol_zona)
        if sel.value in _WARNA_ZONA:
            sel.fill = PatternFill("solid", fgColor=_WARNA_ZONA[sel.value])
    ws.freeze_panes = "B2"
    ws.column_dimensions["A"].width = 22
    for i in range(2, len(header) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 13
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
